"""Google E-Tablolar köprüsü.

disa : Mevcut verileri tablo/72h-getaway-veri.xlsx dosyasına yazar (Google Drive'a yüklenecek tablo) ve
       data/tablo-son.json'a "son görülen tablo" kaydını yazar.
ice  : data/tablo-ayar.json > url adresindeki yayınlanmış tabloyu (Dosya > Paylaş > Web'de yayınla > .xlsx) indirir,
       SADECE son görülen tablodan bu yana değişen hücreleri verilere işler:
       - Oteller   -> data/hotels.json (Google'da aranacak ad değişirse robot oteli yeniden arar)
       - Restoranlar -> data/restoranlar.json (değiştirilen restoran kilitlenir, restoran robotu dokunmaz;
                        Aktif=HAYIR olan robotun bir daha seçmeyeceği listeye girer)
       - Gezi      -> data/gezi.json (aranacak ad değişen ya da yeni yerin konumunu konum robotu bulur)
       - Ulaşım    -> data/gezi.json > ulasim
       Satır silmek hiçbir şeyi silmez; çıkarmak için Aktif=HAYIR yazılır.
Rapor: data/tablo-rapor.json ve (GitHub'da) iş özeti.
"""
import io
import json
import os
import re
import sys
import unicodedata
import urllib.request
from datetime import datetime, timezone

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
P = lambda *a: os.path.join(ROOT, *a)

KATEGORI = ["ikonik", "pahali", "konforlu", "ucuz"]
SEVIYE = ["uygun", "populer", "luks"]
EVET, HAYIR = "EVET", "HAYIR"

SHEETS = {
    "Oteller": {
        "key": ("Şehir", "Kimlik"),
        "cols": ["Şehir", "Kimlik", "Kategori", "Otel adı", "Google'da aranacak ad", "Not", "Aktif"],
        "info": [],
        "widths": [13, 24, 11, 34, 40, 60, 8],
    },
    "Restoranlar": {
        "key": ("Şehir", "Kimlik"),
        "cols": ["Şehir", "Kimlik", "Seviye", "Restoran adı", "Mutfak", "Kişi başı €", "Fiyat yazısı", "Google puanı", "Harita bağlantısı", "Aktif"],
        "info": ["Google puanı"],
        "widths": [13, 30, 9, 36, 24, 11, 18, 11, 40, 8],
    },
    "Gezi": {
        "key": ("Şehir", "Kimlik"),
        "cols": ["Şehir", "Kimlik", "Yer adı", "Tür", "Bölge", "Süre (dk)", "Fiyat", "Hafta sonu fiyatı", "Para birimi", "Fiyat yazısı", "Bilgi", "Kaynak", "Google'da aranacak ad", "Aktif"],
        "info": ["Para birimi"],
        "widths": [13, 28, 32, 13, 12, 9, 9, 10, 9, 40, 60, 40, 36, 8],
    },
    "Ulaşım": {
        "key": ("Şehir",),
        "cols": ["Şehir", "Para birimi", "Tek bilet", "Tek bilet açıklama", "Günlük", "Günlük açıklama", "Çok günlük", "Kaç gün", "Çok günlük açıklama",
                 "Havalimanı", "Havalimanı süre (dk)", "Havalimanı açıklama", "Not", "Kaynak tek bilet", "Kaynak günlük", "Kaynak çok günlük", "Kaynak havalimanı"],
        "info": ["Para birimi"],
        "widths": [13, 9, 9, 40, 9, 40, 10, 8, 40, 10, 10, 40, 50, 30, 30, 30, 30],
    },
}


def load(name, default=None):
    try:
        with open(P(*name.split("/")), encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        if default is not None:
            return default
        raise


def save(name, obj, indent=1):
    with open(P(*name.split("/")), "w", encoding="utf-8") as f:
        f.write(json.dumps(obj, ensure_ascii=False, indent=indent) + "\n")


def slug(s, n=40):
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:n] or "yer"


# ---------- değer normalleştirme (tablo <-> veri) ----------
def norm(v):
    """Karşılaştırma için tek biçim: boş -> '', sayı -> en kısa yazı, metin -> kırpılmış."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return EVET if v else HAYIR
    if isinstance(v, (int, float)):
        f = float(v)
        return str(int(f)) if f.is_integer() else repr(round(f, 4))
    s = str(v).strip()
    if s.upper() in ("EVET", "HAYIR"):
        return s.upper()
    return s


def to_num(v):
    """'2,5' / '2.5' / 2.5 / '' / 'ücretsiz' -> sayı ya da None. Hatalıysa ValueError."""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v) if not float(v).is_integer() else int(v)
    s = str(v).strip().lower().replace("€", "").replace(" ", "")
    if s == "":
        return None
    if s in ("ücretsiz", "ucretsiz", "bedava", "free"):
        return 0
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", s):
        s = s.replace(".", "")
    s = s.replace(",", ".")
    try:
        f = float(s)
    except ValueError:
        raise ValueError(f"sayı değil: {v!r}")
    if f < 0:
        raise ValueError(f"negatif olamaz: {v!r}")
    return int(f) if f.is_integer() else f


def aktif(v):
    s = norm(v).upper()
    return s != HAYIR


# ---------- verilerden tablo satırları ----------
def build_rows():
    hotels = load("data/hotels.json")
    rest = load("data/restoranlar.json", {"sehirler": {}})
    gezi = load("data/gezi.json")
    order = [c["sehir"] for c in hotels["sehirler"]]
    rows = {k: [] for k in SHEETS}
    for c in hotels["sehirler"]:
        for h in c["oteller"]:
            rows["Oteller"].append([c["sehir"], h["id"], h["kategori"], h["ad"], h["arama"], h.get("not", ""), EVET])
    for s in order:
        for r in (rest["sehirler"].get(s) or {}).get("restoranlar", []):
            rows["Restoranlar"].append([s, r["id"], r["seviye"], r["ad"], r.get("mutfak") or "", r.get("kisiBasiEur"), r.get("fiyatYazi") or "", r.get("puan"), r.get("url") or "", EVET])
    for s in order:
        g = gezi["sehirler"].get(s)
        if not g:
            continue
        for y in g["yerler"]:
            rows["Gezi"].append([s, y["id"], y["ad"], y.get("tur") or "", y.get("bolge") or "", y.get("sure_dk"), y.get("fiyat"), y.get("fiyat_hs"), g["para"],
                                 y.get("fiyat_yazi") or "", y.get("bilgi") or "", y.get("kaynak") or "", y.get("arama") or "", EVET])
        u = g.get("ulasim") or {}
        t, d, m, a = (u.get(k) or {} for k in ("tek_bilet", "gunluk", "cok_gunluk", "havalimani"))
        rows["Ulaşım"].append([s, g["para"], t.get("fiyat"), t.get("aciklama", ""), d.get("fiyat"), d.get("aciklama", ""), m.get("fiyat"), m.get("gun"), m.get("aciklama", ""),
                               a.get("fiyat"), a.get("sure_dk"), a.get("aciklama", ""), u.get("not", ""), t.get("kaynak", ""), d.get("kaynak", ""), m.get("kaynak", ""), a.get("kaynak", "")])
    return rows


def snapshot(rows):
    snap = {}
    for name, spec in SHEETS.items():
        cols, kc = spec["cols"], spec["key"]
        out = {}
        for r in rows.get(name, []):
            rec = {c: norm(v) for c, v in zip(cols, r)}
            k = "|".join(rec[c] for c in kc)
            if all(rec[c] for c in kc):
                out[k] = rec
        snap[name] = out
    return snap


# ---------- dışa aktarma ----------
OKUBENI = [
    ("72h Getaway veri tablosu", "baslik"),
    ("Bu tablo sitedeki şehir bilgilerinin düzenleme alanıdır. Robot her sabah tabloyu okur ve SADECE değiştirdiğin hücreleri siteye işler.", ""),
    ("", ""),
    ("Nasıl düzenlerim?", "alt"),
    ("• Bir hücreyi değiştir: ertesi sabah (ya da robotu elle çalıştırınca) sitede görünür.", ""),
    ("• Bir şeyi çıkarmak için satırı SİLME; 'Aktif' sütununa HAYIR yaz. Silinen satır hiçbir şeyi silmez.", ""),
    ("• Yeni otel / restoran / gezi yeri eklemek için en alta yeni satır ekle, 'Kimlik' sütununu BOŞ bırak; robot kimliği kendisi verir.", ""),
    ("• Gri sütunlar bilgi içindir (Google puanı, para birimi); değiştirmen bir şey değiştirmez.", ""),
    ("", ""),
    ("Sekmeler", "alt"),
    ("Oteller: şehir başına seçili oteller. Kategori: ikonik, pahali, konforlu, ucuz. 'Google'da aranacak ad' değişirse robot oteli yeniden arar ve fiyatı ona göre çeker.", ""),
    ("Restoranlar: değiştirdiğin ya da eklediğin restoran KİLİTLENİR; aylık restoran robotu ona dokunmaz. Seviye: uygun, populer, luks. Kişi başı € içecek dahil ortalama.", ""),
    ("Gezi: şehir turu yerleri. Fiyat: yerel para biriminde yetişkin bileti (Para birimi sütununa bak). Ücretsizse 0, bilinmiyorsa boş. Bölge: aynı gün gezilecek yakın yerlere aynı kısa adı ver.", ""),
    ("Ulaşım: toplu taşıma fiyatları (yerel para). Hesaplayıcı havalimanı gidiş-dönüş + en ucuz bilet/kart seçeneğini kullanır.", ""),
    ("", ""),
    ("Sayılar", "alt"),
    ("Ondalık için virgül ya da nokta kullanabilirsin (2,5 ya da 2.5). Para birimi simgesi yazma, sadece sayı.", ""),
    ("", ""),
    ("Yayınlama (bir kez)", "alt"),
    ("Dosya > Paylaş > Web'de yayınla > 'Tüm doküman' ve 'Microsoft Excel (.xlsx)' seç > Yayınla. Çıkan bağlantıyı Claude'a ver.", ""),
]


def export(path=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = build_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "Okubeni"
    ws.column_dimensions["A"].width = 130
    for i, (t, kind) in enumerate(OKUBENI, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "baslik":
            c.font = Font(bold=True, size=16, color="0F2A43")
        elif kind == "alt":
            c.font = Font(bold=True, size=12, color="0F2A43")
    head = PatternFill("solid", fgColor="0F2A43")
    grey = PatternFill("solid", fgColor="EDEDED")
    for name, spec in SHEETS.items():
        sh = wb.create_sheet(name)
        sh.append(spec["cols"])
        for r in rows[name]:
            sh.append(r)
        for j, col in enumerate(spec["cols"], 1):
            cell = sh.cell(row=1, column=j)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sh.column_dimensions[cell.column_letter].width = spec["widths"][j - 1]
            if col in spec["info"] or col == "Kimlik":
                for i in range(2, sh.max_row + 1):
                    sh.cell(row=i, column=j).fill = grey
        sh.freeze_panes = "C2" if name != "Ulaşım" else "B2"
        sh.auto_filter.ref = sh.dimensions
        n = 2000

        def dv(col, values):
            if col not in spec["cols"]:
                return
            j = spec["cols"].index(col) + 1
            letter = sh.cell(row=1, column=j).column_letter
            v = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
            v.add(f"{letter}2:{letter}{n}")
            sh.add_data_validation(v)

        dv("Aktif", [EVET, HAYIR])
        dv("Kategori", KATEGORI)
        dv("Seviye", SEVIYE)
    path = path or P("tablo", "72h-getaway-veri.xlsx")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    save("data/tablo-son.json", {"kaynak": "disa", "tarih": datetime.now(timezone.utc).isoformat(timespec="seconds"), "tablo": snapshot(rows)})
    return path, {k: len(v) for k, v in rows.items()}


# ---------- içe aktarma ----------
def read_xlsx(data):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out, problems = {}, []
    for name, spec in SHEETS.items():
        if name not in wb.sheetnames:
            problems.append(f"'{name}' sekmesi yok")
            continue
        it = wb[name].iter_rows(values_only=True)
        header = [norm(x) for x in next(it, [])]
        missing = [c for c in spec["cols"] if c not in header]
        if missing:
            problems.append(f"'{name}' sekmesinde eksik sütun: {', '.join(missing)}")
            continue
        idx = {c: header.index(c) for c in spec["cols"]}
        rows = []
        for r in it:
            if r is None or all(x in (None, "") for x in r):
                continue
            rows.append({c: (r[idx[c]] if idx[c] < len(r) else None) for c in spec["cols"]})
        out[name] = rows
    return out, problems


def changed_cells(name, row, base):
    """Satırda son görülen tablodan bu yana değişen sütunlar (yeni satırsa None)."""
    spec = SHEETS[name]
    kc = spec["key"]
    k = "|".join(norm(row.get(c)) for c in kc)
    if not all(norm(row.get(c)) for c in kc):
        return None, k
    old = base.get(k)
    if old is None:
        return None, k
    return [c for c in spec["cols"] if c not in kc and c not in spec["info"] and norm(row.get(c)) != old.get(c, "")], k


class Report:
    def __init__(self):
        self.items, self.errors, self.konum = [], [], set()
        self.cur, self.rej = None, set()  # işlenmeyen satırlar: düzeltilene kadar her okumada yeniden raporlanır

    def ok(self, s):
        self.items.append(s)

    def err(self, s):
        self.errors.append(s)
        if self.cur:
            self.rej.add(self.cur)


def rowkey(name, row):
    return "|".join(norm(row.get(c)) for c in SHEETS[name]["key"])


def resolve_ids(sheet, hotels):
    """Kimliği boş (yeni) satırlara addan türeyen sabit kimlik verir; böylece aynı satır her okumada aynı kayda bağlanır."""
    pre = {c["sehir"]: (c["oteller"][0]["id"].split("-")[0] if c["oteller"] else slug(c["sehir"], 3)) for c in hotels["sehirler"]}
    gen = {"Oteller": ("Otel adı", lambda s, ad: f"{pre.get(s, slug(s, 3))}-{slug(ad, 30)}"),
           "Restoranlar": ("Restoran adı", lambda s, ad: "elle-" + slug(ad, 30)),
           "Gezi": ("Yer adı", lambda s, ad: slug(ad))}
    for name, (col, f) in gen.items():
        for row in sheet.get(name, []):
            if not norm(row.get("Kimlik")) and norm(row.get(col)):
                row["Kimlik"] = f(norm(row["Şehir"]), norm(row[col]))
                row["_yeni"] = True


def apply_import(sheet, base, rep):
    hotels = load("data/hotels.json")
    meta = load("data/otel-fiyatlari.json", {})
    rest = load("data/restoranlar.json", {"sehirler": {}})
    gezi = load("data/gezi.json")
    cities = {c["sehir"]: c for c in hotels["sehirler"]}
    touched = set()
    resolve_ids(sheet, hotels)

    # --- Oteller ---
    for row in sheet.get("Oteller", []):
        rep.cur = ("Oteller", rowkey("Oteller", row))
        s = norm(row["Şehir"])
        if s not in cities:
            rep.err(f"Oteller: bilinmeyen şehir '{s}'")
            continue
        city = cities[s]
        ch, key = changed_cells("Oteller", row, base.get("Oteller", {}))
        hid = norm(row["Kimlik"])
        ids = {h["id"] for h in city["oteller"]}
        yeni = row.get("_yeni") and hid not in ids
        if ch == []:
            continue
        if not yeni:
            if hid not in ids:
                if aktif(row["Aktif"]):
                    rep.err(f"Oteller: {s} '{hid}' kimliği bulunamadı (kimliği değiştirme)")
                continue
            if ch is None:
                continue
        kat = norm(row["Kategori"]).lower()
        if kat not in KATEGORI:
            rep.err(f"Oteller: {s} / {norm(row['Otel adı'])}: kategori '{kat}' geçersiz ({', '.join(KATEGORI)})")
            continue
        if yeni:  # yeni otel
            if not aktif(row["Aktif"]):
                continue
            ad, arama = norm(row["Otel adı"]), norm(row["Google'da aranacak ad"]) or (norm(row["Otel adı"]) + " " + s)
            if not ad:
                rep.err(f"Oteller: {s}: yeni satırda otel adı boş")
                continue
            nid = hid
            city["oteller"].append({"id": nid, "ad": ad, "arama": arama, "kategori": kat, "not": norm(row["Not"])})
            rep.ok(f"Otel eklendi: {s} / {ad}")
            touched.add("hotels")
            continue
        h = next(x for x in city["oteller"] if x["id"] == hid)
        if not aktif(row["Aktif"]):
            city["oteller"].remove(h)
            rep.ok(f"Otel çıkarıldı: {s} / {h['ad']}")
            touched.add("hotels")
            continue
        fields = {"Kategori": "kategori", "Otel adı": "ad", "Google'da aranacak ad": "arama", "Not": "not"}
        for c in ch:
            if c in fields:
                h[fields[c]] = kat if c == "Kategori" else norm(row[c])
        if "Google'da aranacak ad" in ch:
            (meta.get("googleKimlik") or {}).pop(hid, None)
            (meta.get("konum") or {}).pop(hid, None)
        rep.ok(f"Otel güncellendi: {s} / {h['ad']} ({', '.join(ch)})")
        touched.add("hotels")

    # --- Restoranlar ---
    for row in sheet.get("Restoranlar", []):
        rep.cur = ("Restoranlar", rowkey("Restoranlar", row))
        s = norm(row["Şehir"])
        if s not in cities:
            rep.err(f"Restoranlar: bilinmeyen şehir '{s}'")
            continue
        sc = rest["sehirler"].setdefault(s, {"restoranlar": []})
        lst = sc.setdefault("restoranlar", [])
        ch, key = changed_cells("Restoranlar", row, base.get("Restoranlar", {}))
        rid = norm(row["Kimlik"])
        yeni = row.get("_yeni") and not any(x["id"] == rid for x in lst) and rid not in sc.get("haric", [])
        if ch == []:
            continue
        if ch is None and not yeni:
            continue  # tabloda var, son görülende yok: dokunma (robotun eklediği eski satır olabilir)
        lvl = norm(row["Seviye"]).lower()
        if lvl not in SEVIYE:
            rep.err(f"Restoranlar: {s} / {norm(row['Restoran adı'])}: seviye '{lvl}' geçersiz ({', '.join(SEVIYE)})")
            continue
        try:
            eur = to_num(row["Kişi başı €"])
        except ValueError as e:
            rep.err(f"Restoranlar: {s} / {norm(row['Restoran adı'])}: kişi başı € {e}")
            continue
        url = norm(row["Harita bağlantısı"])
        if url and not url.startswith("https://"):
            rep.err(f"Restoranlar: {s} / {norm(row['Restoran adı'])}: harita bağlantısı https:// ile başlamalı")
            continue
        if yeni:  # yeni restoran
            if not aktif(row["Aktif"]):
                continue
            ad = norm(row["Restoran adı"])
            if not ad or eur is None:
                rep.err(f"Restoranlar: {s}: yeni satırda ad ve kişi başı € zorunlu")
                continue
            nid = rid
            lst.append({"id": nid, "ad": ad, "mutfak": norm(row["Mutfak"]) or None, "puan": None, "yorum": 0, "fiyatYazi": norm(row["Fiyat yazısı"]) or f"~{eur} €",
                        "kisiBasiEur": eur, "artiFiyat": False, "semt": None, "konum": None, "url": url or None, "web": None, "rezervasyon": None, "seviye": lvl, "kilit": True})
            rep.ok(f"Restoran eklendi (kilitli): {s} / {ad}")
            touched.add("rest")
            continue
        r = next((x for x in lst if x["id"] == rid), None)
        if not aktif(row["Aktif"]):
            if r:
                lst.remove(r)
            hx = sc.setdefault("haric", [])
            if rid not in hx:
                hx.append(rid)
            rep.ok(f"Restoran çıkarıldı, robot bir daha seçmeyecek: {s} / {norm(row['Restoran adı'])}")
            touched.add("rest")
            continue
        if r is None:  # robot çıkarmış ama kullanıcı düzenlemiş: kullanıcının değerleriyle geri ekle
            r = {"id": rid, "puan": None, "yorum": 0, "artiFiyat": False, "semt": None, "konum": None, "web": None, "rezervasyon": None}
            lst.append(r)
        r.update({"ad": norm(row["Restoran adı"]) or r.get("ad"), "mutfak": norm(row["Mutfak"]) or r.get("mutfak"), "seviye": lvl,
                  "kisiBasiEur": eur if eur is not None else r.get("kisiBasiEur"), "fiyatYazi": norm(row["Fiyat yazısı"]) or r.get("fiyatYazi"),
                  "url": url or r.get("url"), "kilit": True})
        r.pop("yedek", None)
        rep.ok(f"Restoran güncellendi (kilitli): {s} / {r['ad']} ({', '.join(ch or [])})")
        touched.add("rest")

    # --- Gezi ---
    for row in sheet.get("Gezi", []):
        rep.cur = ("Gezi", rowkey("Gezi", row))
        s = norm(row["Şehir"])
        g = gezi["sehirler"].get(s)
        if not g:
            rep.err(f"Gezi: bilinmeyen şehir '{s}'")
            continue
        ch, key = changed_cells("Gezi", row, base.get("Gezi", {}))
        yid = norm(row["Kimlik"])
        yeni = row.get("_yeni") and not any(y["id"] == yid for y in g["yerler"])
        if ch == [] or (ch is None and not yeni):
            continue
        try:
            fiyat, fhs, sure = to_num(row["Fiyat"]), to_num(row["Hafta sonu fiyatı"]), to_num(row["Süre (dk)"])
        except ValueError as e:
            rep.err(f"Gezi: {s} / {norm(row['Yer adı'])}: {e}")
            continue
        if sure is not None and not (10 <= sure <= 480):
            rep.err(f"Gezi: {s} / {norm(row['Yer adı'])}: süre 10–480 dk olmalı")
            continue
        vals = {"ad": norm(row["Yer adı"]), "tur": norm(row["Tür"]), "bolge": slug(norm(row["Bölge"]), 20) if norm(row["Bölge"]) else "", "sure_dk": sure, "fiyat": fiyat, "fiyat_hs": fhs,
                "fiyat_yazi": norm(row["Fiyat yazısı"]), "bilgi": norm(row["Bilgi"]), "kaynak": norm(row["Kaynak"]), "arama": norm(row["Google'da aranacak ad"])}
        if yeni:  # yeni yer
            if not aktif(row["Aktif"]):
                continue
            if not vals["ad"] or not sure:
                rep.err(f"Gezi: {s}: yeni satırda yer adı ve süre zorunlu")
                continue
            nid = yid
            y = {"id": nid, **vals, "konum": None}
            y["arama"] = y["arama"] or f"{vals['ad']} {s}"
            y["tur"] = y["tur"] or "Gezi"
            y["bolge"] = y["bolge"] or nid[:12]
            y["fiyat_yazi"] = y["fiyat_yazi"] or ("Ücretsiz" if fiyat == 0 else "Fiyatı teyit et")
            y["bilgi"] = y["bilgi"] or vals["ad"]
            g["yerler"].append(y)
            rep.konum.add(s)
            rep.ok(f"Gezi yeri eklendi: {s} / {vals['ad']} (konumu robot bulacak)")
            touched.add("gezi")
            continue
        y = next((x for x in g["yerler"] if x["id"] == yid), None)
        if y is None:
            rep.err(f"Gezi: {s} '{yid}' kimliği bulunamadı (kimliği değiştirme)")
            continue
        if not aktif(row["Aktif"]):
            g["yerler"].remove(y)
            rep.ok(f"Gezi yeri çıkarıldı: {s} / {y['ad']}")
            touched.add("gezi")
            continue
        colmap = {"Yer adı": "ad", "Tür": "tur", "Bölge": "bolge", "Süre (dk)": "sure_dk", "Fiyat": "fiyat", "Hafta sonu fiyatı": "fiyat_hs", "Fiyat yazısı": "fiyat_yazi", "Bilgi": "bilgi", "Kaynak": "kaynak", "Google'da aranacak ad": "arama"}
        for c in ch:
            if c in colmap:
                v = vals[colmap[c]]
                if colmap[c] in ("ad", "sure_dk", "tur", "bilgi", "fiyat_yazi", "arama") and v in ("", None):
                    rep.err(f"Gezi: {s} / {y['ad']}: '{c}' boş bırakılamaz")
                    continue
                y[colmap[c]] = v
        if "Google'da aranacak ad" in ch:
            y["konum"] = None
            for k in ("googleAd", "konumNotu", "puan", "yorum", "harita", "kapali"):
                y.pop(k, None)
            rep.konum.add(s)
        rep.ok(f"Gezi yeri güncellendi: {s} / {y['ad']} ({', '.join(ch)})")
        touched.add("gezi")

    # --- Ulaşım ---
    for row in sheet.get("Ulaşım", []):
        rep.cur = ("Ulaşım", rowkey("Ulaşım", row))
        s = norm(row["Şehir"])
        g = gezi["sehirler"].get(s)
        if not g:
            rep.err(f"Ulaşım: bilinmeyen şehir '{s}'")
            continue
        ch, key = changed_cells("Ulaşım", row, base.get("Ulaşım", {}))
        if not ch:
            continue
        u = g.setdefault("ulasim", {})
        try:
            nums = {c: to_num(row[c]) for c in ("Tek bilet", "Günlük", "Çok günlük", "Kaç gün", "Havalimanı", "Havalimanı süre (dk)")}
        except ValueError as e:
            rep.err(f"Ulaşım: {s}: {e}")
            continue
        spec = [("tek_bilet", "Tek bilet", "Tek bilet açıklama", "Kaynak tek bilet"), ("gunluk", "Günlük", "Günlük açıklama", "Kaynak günlük"),
                ("cok_gunluk", "Çok günlük", "Çok günlük açıklama", "Kaynak çok günlük"), ("havalimani", "Havalimanı", "Havalimanı açıklama", "Kaynak havalimanı")]
        for k, cp, ca, ck in spec:
            if not ({cp, ca, ck} & set(ch)) and not (k == "cok_gunluk" and "Kaç gün" in ch) and not (k == "havalimani" and "Havalimanı süre (dk)" in ch):
                continue
            if nums[cp] is None:
                if k in ("tek_bilet", "havalimani"):
                    rep.err(f"Ulaşım: {s}: '{cp}' boş bırakılamaz")
                    continue
                u[k] = None
                continue
            o = dict(u.get(k) or {})
            o.update({"fiyat": nums[cp], "aciklama": norm(row[ca]), "kaynak": norm(row[ck])})
            if k == "cok_gunluk":
                o["gun"] = int(nums["Kaç gün"] or 3)
            if k == "havalimani" and nums["Havalimanı süre (dk)"] is not None:
                o["sure_dk"] = nums["Havalimanı süre (dk)"]
            u[k] = o
        if "Not" in ch:
            u["not"] = norm(row["Not"])
        rep.ok(f"Ulaşım güncellendi: {s} ({', '.join(ch)})")
        touched.add("gezi")

    # kaydet
    if "hotels" in touched:
        for c in hotels["sehirler"]:
            if not c["oteller"]:
                raise SystemExit(f"Güvenlik: {c['sehir']} için hiç otel kalmıyor, aktarma durduruldu")
        save_hotels(hotels)
        if meta:
            save("data/otel-fiyatlari.json", meta)
    if "rest" in touched:
        save("data/restoranlar.json", rest)
    if "gezi" in touched:
        gezi["guncellendi"] = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        save("data/gezi.json", gezi)
    return touched


def save_hotels(doc):
    """hotels.json'u okunaklı biçimde (otel başına bir satır) yazar."""
    J = lambda o: json.dumps(o, ensure_ascii=False, separators=(",", ":"))
    a = doc["ayarlar"]
    lines = ["{", f'  "aciklama": {json.dumps(doc.get("aciklama", ""), ensure_ascii=False)},', '  "ayarlar": {']
    keys = [k for k in a if k != "kontrol_takvimi"]
    lines += [f'    {json.dumps(k)}: {J(a[k])},' for k in keys]
    lines += ['    "kontrol_takvimi": [', ",\n".join("      " + J(t) for t in a["kontrol_takvimi"]), "    ]", "  },", '  "sehirler": [']
    cs = []
    for c in doc["sehirler"]:
        cs.append("    {\n" + f'      "sehir": {json.dumps(c["sehir"], ensure_ascii=False)},\n      "ulke": {json.dumps(c["ulke"], ensure_ascii=False)},\n      "oteller": [\n'
                  + ",\n".join("        " + J(h) for h in c["oteller"]) + "\n      ]\n    }")
    lines += [",\n".join(cs), "  ]", "}"]
    text = "\n".join(lines) + "\n"
    json.loads(text)
    with open(P("data", "hotels.json"), "w", encoding="utf-8") as f:
        f.write(text)


def import_(data_bytes):
    rep = Report()
    sheet, problems = read_xlsx(data_bytes)
    if problems:
        rep.errors += problems
        return rep, set(), False
    base = (load("data/tablo-son.json", {}).get("tablo")) or {}
    touched = apply_import(sheet, base, rep)
    # yeni "son görülen tablo": şimdi okunan tablo
    rows = {name: [[r.get(c) for c in SHEETS[name]["cols"]] for r in sheet.get(name, [])] for name in SHEETS}
    snap = snapshot(rows)
    for name, k in rep.rej:  # işlenmeyen satır: eski hâli kalsın ki düzeltilince değişiklik olarak görülsün
        if k in (base.get(name) or {}):
            snap[name][k] = base[name][k]
        else:
            snap[name].pop(k, None)
    if snap != base:  # tablo değişmediyse dosyaya dokunma (boş yere kayıt olmasın)
        save("data/tablo-son.json", {"kaynak": "ice", "tarih": datetime.now(timezone.utc).isoformat(timespec="seconds"), "tablo": snap})
    return rep, touched, True


def download_urls(url):
    """Kabul edilen bağlantılar: 'Web'de yayınla' (.../d/e/.../pub?output=xlsx) ya da tablonun/dosyanın kendi bağlantısı
    (docs.google.com/spreadsheets/d/<kimlik>/..., drive.google.com/file/d/<kimlik>/...). Kendi bağlantısında dosya
    'Bağlantıya sahip olan herkes görüntüleyebilir' olarak paylaşılmış olmalı."""
    if re.match(r"^https://docs\.google\.com/spreadsheets/d/e/[\w-]+/pub\?output=xlsx$", url):
        return [url]
    m = re.match(r"^https://(?:docs\.google\.com/spreadsheets|drive\.google\.com/file)/d/([\w-]{20,})", url)
    if not m:
        raise SystemExit("Tablo bağlantısı tanınmadı: Google E-Tablolar ya da Drive bağlantısı olmalı")
    i = m.group(1)
    return [f"https://docs.google.com/spreadsheets/d/{i}/export?format=xlsx", f"https://drive.google.com/uc?export=download&id={i}"]


def main():
    mode = (sys.argv[1] if len(sys.argv) > 1 else "ice").strip()
    if mode == "disa":
        path, counts = export()
        print("Tablo yazıldı:", path, counts)
        return
    ayar = load("data/tablo-ayar.json", {"url": ""})
    url = (ayar.get("url") or "").strip()
    if not url:
        print("Tablo bağlantısı yok (data/tablo-ayar.json > url). Aktarma atlandı.")
        return
    data, tried = None, []
    for u in download_urls(url):
        try:
            req = urllib.request.Request(u, headers={"User-Agent": "72h-getaway-tablo-robotu"})
            with urllib.request.urlopen(req, timeout=60) as r:
                body = r.read()
        except Exception as e:  # noqa: BLE001
            tried.append(f"{u.split('?')[0]}: {e}")
            continue
        if body.startswith(b"PK"):
            data = body
            break
        tried.append(f"{u.split('?')[0]}: xlsx gelmedi (giriş sayfası)")
    if data is None:
        save("data/tablo-rapor.json", {"tarih": datetime.now(timezone.utc).isoformat(timespec="seconds"), "islenen": [], "hatalar": ["Tablo indirilemedi: dosya 'Bağlantıya sahip olan herkes' ile paylaşılmamış olabilir", *tried], "konum_aranacak": []})
        raise SystemExit("Tablo indirilemedi. Google'da Paylaş > Genel erişim > 'Bağlantıya sahip olan herkes' (Görüntüleyen) seçili olmalı.\n" + "\n".join(tried))
    rep, touched, ok = import_(data)
    md = ["# Tablo aktarımı", "", f"Değişiklik: {len(rep.items)}, hata: {len(rep.errors)}.", ""]
    md += ["## İşlenenler", *[f"- {x}" for x in rep.items], ""] if rep.items else ["Tabloda yeni değişiklik yok.", ""]
    md += ["## Düzeltilmesi gerekenler (işlenmedi)", *[f"- {x}" for x in rep.errors], ""] if rep.errors else []
    text = "\n".join(md)
    print(text)
    save("data/tablo-rapor.json", {"tarih": datetime.now(timezone.utc).isoformat(timespec="seconds"), "islenen": rep.items, "hatalar": rep.errors, "konum_aranacak": sorted(rep.konum)})
    if os.environ.get("GITHUB_STEP_SUMMARY"):
        with open(os.environ["GITHUB_STEP_SUMMARY"], "a", encoding="utf-8") as f:
            f.write(text + "\n")
    if os.environ.get("GITHUB_OUTPUT"):
        with open(os.environ["GITHUB_OUTPUT"], "a", encoding="utf-8") as f:
            f.write("konum=" + ",".join(sorted(rep.konum)) + "\n")
    if not ok:
        raise SystemExit("Tablo yapısı bozuk, hiçbir şey aktarılmadı: " + "; ".join(rep.errors))


if __name__ == "__main__":
    main()
