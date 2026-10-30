"""Tablo köprüsünün testleri. Çalıştır: python3 scripts/test_tablo.py
Depoyu geçici klasöre kopyalar; gerçek veriye dokunmaz."""
import importlib.util
import io
import json
import os
import shutil
import sys
import tempfile

SRC = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
tmp = tempfile.mkdtemp()
for d in ("scripts", "data"):
    shutil.copytree(os.path.join(SRC, d), os.path.join(tmp, d))
spec = importlib.util.spec_from_file_location("tablo", os.path.join(tmp, "scripts", "tablo.py"))
T = importlib.util.module_from_spec(spec)
spec.loader.exec_module(T)

from openpyxl import load_workbook

passed = 0


def check(name, cond, extra=""):
    global passed
    if not cond:
        print("BAŞARISIZ:", name, extra)
        sys.exit(1)
    passed += 1
    print("ok", name)


def read(name):
    with open(os.path.join(tmp, "data", name), encoding="utf-8") as f:
        return f.read()


def wb_bytes(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def find(ws, **match):
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if all(row[head.index(k)].value == v for k, v in match.items()):
            return {h: row[i] for i, h in enumerate(head)}
    raise KeyError(match)


# 1. dışa aktarma
path, counts = T.export(os.path.join(tmp, "t.xlsx"))
check("1 sekme sayıları", counts["Oteller"] > 100 and counts["Ulaşım"] == 26, counts)
before = {n: read(n) for n in ("hotels.json", "restoranlar.json", "gezi.json")}

# 2. değişmemiş tablo hiçbir şeyi değiştirmez
with open(path, "rb") as f:
    rep, touched, ok = T.import_(f.read())
check("2 değişiklik yok", ok and not rep.items and not rep.errors and not touched, (rep.items, rep.errors))
check("2b dosyalar aynı", all(read(n) == before[n] for n in before))

# 3. hotels.json biçimi korunuyor (yeniden yazınca aynı metin)
T.save_hotels(json.loads(before["hotels.json"]))
check("3 hotels.json biçimi", read("hotels.json") == before["hotels.json"])

# 4. değişiklikler
wb = load_workbook(path)
g = wb["Gezi"]
hof = find(g, **{"Şehir": "Viyana"})  # ilk Viyana yeri
hof_id = hof["Kimlik"].value
hof["Fiyat"].value = "24,5"
hof["Fiyat yazısı"].value = "Yetişkin 24,50 €"
g.append(["Viyana", None, "Test Müzesi", "Müze", "merkez", 60, 0, None, "EUR", "", "Deneme", "", "", "EVET"])
r = wb["Restoranlar"]
rr = find(r, **{"Şehir": "Paris"})
r_id = rr["Kimlik"].value
rr["Kişi başı €"].value = 33
rr2 = None
for row in r.iter_rows(min_row=2):
    if row[0].value == "Paris" and row[1].value != r_id:
        rr2 = row
        break
r_off = rr2[1].value
rr2[9].value = "HAYIR"
o = wb["Oteller"]
oo = find(o, **{"Şehir": "Roma"})
o_id = oo["Kimlik"].value
oo["Google'da aranacak ad"].value = "Hotel De Russie Rome Rocco Forte"
u = wb["Ulaşım"]
uu = find(u, **{"Şehir": "Prag"})
uu["Tek bilet"].value = 31
bad = find(o, **{"Şehir": "Madrid"})
bad["Kategori"].value = "süper"
rep, touched, ok = T.import_(wb_bytes(wb))
check("4 işlendi", ok and len(rep.items) == 6, rep.items)
check("4b hatalı kategori raporlandı", len(rep.errors) == 1 and "kategori" in rep.errors[0], rep.errors)
gz = json.loads(read("gezi.json"))
y = next(y for y in gz["sehirler"]["Viyana"]["yerler"] if y["id"] == hof_id)
check("4c gezi fiyatı virgüllü sayı", y["fiyat"] == 24.5 and y["fiyat_yazi"] == "Yetişkin 24,50 €")
yeni = next(y for y in gz["sehirler"]["Viyana"]["yerler"] if y["ad"] == "Test Müzesi")
check("4d yeni yer kimlik + konum aranacak", yeni["id"] == "test-muzesi" and yeni["konum"] is None and "Viyana" in rep.konum and yeni["fiyat_yazi"] == "Ücretsiz")
check("4e ulaşım", gz["sehirler"]["Prag"]["ulasim"]["tek_bilet"]["fiyat"] == 31)
rs = json.loads(read("restoranlar.json"))["sehirler"]["Paris"]
x = next(x for x in rs["restoranlar"] if x["id"] == r_id)
check("4f restoran kilitli ve güncel", x["kisiBasiEur"] == 33 and x["kilit"] is True)
check("4g çıkarılan restoran hariç listesinde", r_off in rs["haric"] and not any(x["id"] == r_off for x in rs["restoranlar"]))
ho = json.loads(read("hotels.json"))
h = next(h for c in ho["sehirler"] for h in c["oteller"] if h["id"] == o_id)
check("4h otel arama adı", h["arama"] == "Hotel De Russie Rome Rocco Forte")
mad = next(c for c in ho["sehirler"] if c["sehir"] == "Madrid")
check("4i hatalı satır işlenmedi", all(h["kategori"] in T.KATEGORI for h in mad["oteller"]))

# 5. aynı tablo ikinci kez: yeni değişiklik yok (yeni satır tekrar eklenmez)
rep2, touched2, ok2 = T.import_(wb_bytes(wb))
check("5 ikinci okuma boş", ok2 and not rep2.items, rep2.items)
check("5c düzeltilmeyen hata yine raporlanır", len(rep2.errors) == 1, rep2.errors)
bad2 = find(wb["Oteller"], **{"Şehir": "Madrid", "Kategori": "süper"})
bad2["Kategori"].value = "ucuz"
rep4, _, _ = T.import_(wb_bytes(wb))
check("5d düzeltilince işlenir", len(rep4.items) == 1 and not rep4.errors, (rep4.items, rep4.errors))
gz2 = json.loads(read("gezi.json"))
check("5b yeni yer tek kez", sum(1 for y in gz2["sehirler"]["Viyana"]["yerler"] if y["ad"] == "Test Müzesi") == 1)

# 6. bozuk yapı hiçbir şey aktarmaz
wb3 = load_workbook(path)
del wb3["Gezi"]
snap_before = read("tablo-son.json")
rep3, t3, ok3 = T.import_(wb_bytes(wb3))
check("6 eksik sekme durdurur", not ok3 and "Gezi" in rep3.errors[0] and read("tablo-son.json") == snap_before)

# 7. sayı okuma
check("7 sayılar", T.to_num("1.250,5") == 1250.5 and T.to_num("ücretsiz") == 0 and T.to_num("") is None and T.to_num(" 12 € ") == 12)
try:
    T.to_num("on beş")
    check("7b hatalı sayı", False)
except ValueError:
    check("7b hatalı sayı", True)

# 8. restoran robotu kilitli/haric kuralına uyuyor mu: node testinde

shutil.rmtree(tmp)
print(f"\n{passed} test geçti")
